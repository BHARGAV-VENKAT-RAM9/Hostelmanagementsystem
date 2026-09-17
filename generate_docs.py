import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document(output_path):
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Color Palette Constants
    PRIMARY = RGBColor(30, 41, 59)     # Slate Navy (#1E293B)
    SECONDARY = RGBColor(14, 116, 144) # Ocean Teal (#0E7490)
    DARK_TEXT = RGBColor(51, 65, 85)   # Body Text (#334155)
    MUTED = RGBColor(100, 116, 139)    # Muted (#64748B)

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = DARK_TEXT

    # ----------------------------------------------------
    # TITLE SECTION
    # ----------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(20)
    title_p.paragraph_format.space_after = Pt(6)
    
    title_run = title_p.add_run("LUXEHOSTEL MANAGEMENT SYSTEM")
    title_run.bold = True
    title_run.font.size = Pt(24)
    title_run.font.color.rgb = PRIMARY

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(16)
    sub_run = subtitle_p.add_run("Comprehensive Project Documentation & System Specifications")
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = SECONDARY

    # Meta banner / horizontal box
    meta_table = doc.add_table(rows=1, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_row = meta_table.rows[0]
    meta_data = [
        ("Author / Owner", "BHARGAV-VENKAT-RAM9"),
        ("Technology Stack", "React 19, Node.js, Express"),
        ("Repository", "Hostelmanagementsystem")
    ]
    for i, (label, val) in enumerate(meta_data):
        cell = meta_row.cells[i]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{label}\n")
        r1.font.size = Pt(9.5)
        r1.font.bold = True
        r1.font.color.rgb = SECONDARY
        r2 = p.add_run(val)
        r2.font.size = Pt(10)
        r2.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ----------------------------------------------------
    # Helper to add Headings
    # ----------------------------------------------------
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = PRIMARY
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = SECONDARY
        return h

    def add_bullet(p, bold_prefix, text):
        p.style = 'List Bullet'
        p.paragraph_format.space_after = Pt(3)
        r_bold = p.add_run(bold_prefix)
        r_bold.bold = True
        r_bold.font.color.rgb = PRIMARY
        r_text = p.add_run(text)
        r_text.font.color.rgb = DARK_TEXT

    # ----------------------------------------------------
    # SECTION 1: EXECUTIVE SUMMARY
    # ----------------------------------------------------
    add_heading_1("1. Executive Summary")
    p = doc.add_paragraph(
        "LuxeHostel is an end-to-end, full-stack hostel and accommodation management web application designed "
        "to modernise educational, corporate, and private hostel operations. Traditional hostel workflows frequently "
        "suffer from fragmented paper logs, delayed rent accounting, physical complaint submission bottlenecks, and "
        "opaque vacancy tracking. LuxeHostel addresses these hurdles by uniting prospective visitors, student residents, "
        "and hostel administrators under a high-performance, real-time digital architecture."
    )
    p.paragraph_format.space_after = Pt(8)

    p = doc.add_paragraph(
        "The application is engineered with an isolated two-tier frontend architecture—providing an engaging, public "
        "floor grid and resident portal on one client tier, complemented by an administrative command center running "
        "on a dedicated port. Behind the scenes, a lightweight, highly responsive Node.js/Express REST API coordinates "
        "instant vacancy recalculations, check-ins/check-outs, recurring monthly billing, and maintenance ticketing."
    )
    p.paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 2: PROBLEM STATEMENT & MOTIVATION
    # ----------------------------------------------------
    add_heading_1("2. Problem Statement & Motivation")
    p = doc.add_paragraph(
        "Managing residential hostels encompasses numerous moving parts that create friction when handled manually:"
    )
    p.paragraph_format.space_after = Pt(4)

    b1 = doc.add_paragraph()
    add_bullet(b1, "Manual Ledger Inaccuracies: ", "Recording bed allocations, vacancies, and occupant contact details manually often leads to double-booking and data discrepancies.")
    b2 = doc.add_paragraph()
    add_bullet(b2, "Cumbersome Fee Collection: ", "Tracking monthly hostel dues, managing mixed payment modes (cash vs. digital UPI), and printing receipts requires extensive administrative effort.")
    b3 = doc.add_paragraph()
    add_bullet(b3, "Lack of Real-Time Vacancy Visibility: ", "Prospective occupants and parents cannot ascertain which rooms (AC or Non-AC) or specific beds are vacant without physically visiting the premises.")
    b4 = doc.add_paragraph()
    add_bullet(b4, "Untracked Maintenance Complaints: ", "Verbal repair requests for plumbing, electrical, or Wi-Fi issues often go misplaced or unresolved due to a lack of accountable audit trails.")

    # ----------------------------------------------------
    # SECTION 3: SYSTEM ARCHITECTURE & TECH STACK
    # ----------------------------------------------------
    add_heading_1("3. System Architecture & Technology Stack")
    p = doc.add_paragraph(
        "The project follows a modern Decoupled Client-Server Architecture adhering to RESTful design patterns:"
    )
    p.paragraph_format.space_after = Pt(6)

    # Table of Tech Stack
    tech_table = doc.add_table(rows=6, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer / Component", "Technology / Framework", "Role & Description"]
    
    # Format Header Row
    for i, h_text in enumerate(headers):
        cell = tech_table.rows[0].cells[i]
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        ("Frontend (Client)", "React.js 19 + Vite 8", "Delivers reactive UI components for Public Floor Grid and Student Dashboard."),
        ("Admin Portal", "React.js 19 + Vite 8", "Dedicated, secure administrative interface for hostel managers and wardens."),
        ("Styling & Design", "Modern CSS3 + Glassmorphism", "Custom design system featuring CSS variables, responsive grids, and dark/glow styling."),
        ("Backend Server", "Node.js + Express.js", "Handles business logic, REST endpoints, CORS policies, and state persistence."),
        ("Persistence Layer", "MongoDB / Persistent JSON DB", "Atomic read/write operations for rooms, beds, billing history, and complaints.")
    ]

    for row_idx, data in enumerate(rows_data, start=1):
        for col_idx, text in enumerate(data):
            cell = tech_table.rows[row_idx].cells[col_idx]
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 4: CORE MODULES & FEATURES
    # ----------------------------------------------------
    add_heading_1("4. Detailed Module Breakdown")

    add_heading_2("4.1. Public & Visitor Floor Grid Explorer")
    p = doc.add_paragraph(
        "Designed to provide transparent, unauthenticated discovery for visitors and prospective residents:"
    )
    p.paragraph_format.space_after = Pt(4)
    v1 = doc.add_paragraph()
    add_bullet(v1, "5-Floor Visual Grid: ", "Renders all 5 floors and 50 rooms organized by category (Floors 1-3 Non-AC, Floors 4-5 AC).")
    v2 = doc.add_paragraph()
    add_bullet(v2, "Occupancy & Bed Badging: ", "Interactive bed status indicator (Occupied vs. Available) with color-coded status chips.")
    v3 = doc.add_paragraph()
    add_bullet(v3, "Summary Statistics: ", "Real-time cards computing Total Beds, Occupied Beds, and Overall Occupancy Rate.")

    add_heading_2("4.2. Student Resident Portal")
    p = doc.add_paragraph(
        "Personalized resident dashboard activated upon student login:"
    )
    p.paragraph_format.space_after = Pt(4)
    s1 = doc.add_paragraph()
    add_bullet(s1, "Room & Bed Overview: ", "Displays assigned room number, bed slot (e.g., Bed A), room amenity tier, and monthly rate.")
    s2 = doc.add_paragraph()
    add_bullet(s2, "LuxeCoins Loyalty Rewards: ", "Students accumulate loyalty coins from prompt payments that can be redeemed for rent discounts.")
    s3 = doc.add_paragraph()
    add_bullet(s3, "Maintenance Helpdesk: ", "Allows students to lodge tickets categorized by Cleaning, Electrical, Plumbing, Wi-Fi, or Furniture.")

    add_heading_2("4.3. Interactive Payment Gateway & Invoicing")
    p = doc.add_paragraph(
        "A multi-modal payment system simulating complete transaction workflows:"
    )
    p.paragraph_format.space_after = Pt(4)
    pg1 = doc.add_paragraph()
    add_bullet(pg1, "Payment Modes: ", "Supports simulated UPI (VPA address), Credit/Debit Card (card formatting & expiry validation), and Net Banking.")
    pg2 = doc.add_paragraph()
    add_bullet(pg2, "Flexible Settlement: ", "Allows students to pay 100% online, record cash payments, or perform split transactions.")
    pg3 = doc.add_paragraph()
    add_bullet(pg3, "Printable Digital Receipts: ", "Generates structured, professional receipts with unique Transaction IDs, timestamps, and breakdown.")

    add_heading_2("4.4. Dedicated Administrator Command Center")
    p = doc.add_paragraph(
        "A management suite for hostel authorities hosted independently on Port 5174:"
    )
    p.paragraph_format.space_after = Pt(4)
    a1 = doc.add_paragraph()
    add_bullet(a1, "Bed Allocation & Check-In: ", "Seamless student check-in capturing full name, mobile number, and auto-generated credentials.")
    a2 = doc.add_paragraph()
    add_bullet(a2, "Immediate Check-Out: ", "One-click vacate functionality that clears occupant records and returns the bed to available inventory.")
    a3 = doc.add_paragraph()
    add_bullet(a3, "Automated Monthly Billing Engine: ", "Batch invoice generator calculating upcoming dues across all active residents with month-wise filtering.")
    a4 = doc.add_paragraph()
    add_bullet(a4, "Dynamic Rate Configuration: ", "Allows admins to update base pricing for AC and Non-AC rooms with immediate system-wide sync.")
    a5 = doc.add_paragraph()
    add_bullet(a5, "Complaint Resolution Workflow: ", "Centralised queue enabling admins to review complaints, examine timestamps, and update statuses to Resolved.")

    # ----------------------------------------------------
    # SECTION 5: REST API SPECIFICATION
    # ----------------------------------------------------
    add_heading_1("5. REST API Specifications")
    p = doc.add_paragraph(
        "The Express.js backend exposes clear RESTful endpoints for all state queries and mutations:"
    )
    p.paragraph_format.space_after = Pt(6)

    api_table = doc.add_table(rows=8, cols=3)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    api_headers = ["Endpoint", "HTTP Method", "Functional Purpose"]
    
    for i, h_text in enumerate(api_headers):
        cell = api_table.rows[0].cells[i]
        set_cell_background(cell, "0E7490")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    api_data = [
        ("/api/state", "GET", "Fetches full system state: rooms, beds, rates, complaints, and bills."),
        ("/api/login", "POST", "Authenticates student or administrator accounts."),
        ("/api/checkin", "POST", "Allocates a specific bed to a new student and sets rent status."),
        ("/api/checkout", "POST", "Vacates an occupied bed and resets room occupancy."),
        ("/api/pay-bill", "POST", "Processes payment for a bed or bill invoice and issues a receipt."),
        ("/api/generate-bills", "POST", "Generates recurring monthly rent bills for all occupied beds."),
        ("/api/complaints", "POST / PUT", "Creates a new repair ticket or updates ticket resolution status.")
    ]

    for row_idx, data in enumerate(api_data, start=1):
        for col_idx, text in enumerate(data):
            cell = api_table.rows[row_idx].cells[col_idx]
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ----------------------------------------------------
    # SECTION 6: INSTALLATION & EXECUTION GUIDE
    # ----------------------------------------------------
    add_heading_1("6. Installation & Execution Guide")
    p = doc.add_paragraph(
        "To launch the full suite locally, open three separate terminal windows and run the following commands:"
    )
    p.paragraph_format.space_after = Pt(6)

    def add_code_block(title, code_lines):
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(4)
        cp.paragraph_format.space_after = Pt(2)
        r_title = cp.add_run(title)
        r_title.bold = True
        r_title.font.color.rgb = PRIMARY
        
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for line in code_lines:
            r = p.add_run(line + "\n")
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(30, 41, 59)

    add_code_block("Terminal 1 — Backend REST API Server (Port 5000):", [
        "cd backend",
        "npm install",
        "npm run dev"
    ])

    add_code_block("Terminal 2 — Student & Visitor Client Portal (Port 5173):", [
        "cd frontend",
        "npm install",
        "npm run dev"
    ])

    add_code_block("Terminal 3 — Admin Management Portal (Port 5174):", [
        "cd adminfrontend",
        "npm install",
        "npm run dev"
    ])

    # ----------------------------------------------------
    # SECTION 7: CONCLUSION & FUTURE ENHANCEMENTS
    # ----------------------------------------------------
    add_heading_1("7. Conclusion & Future Roadmap")
    p = doc.add_paragraph(
        "The LuxeHostel Management System provides a solid foundation for digital hostel administration. "
        "Planned enhancements for future iterations include:"
    )
    p.paragraph_format.space_after = Pt(4)
    f1 = doc.add_paragraph()
    add_bullet(f1, "Live Payment Gateway Webhooks: ", "Integrating production Razorpay / Stripe webhooks for real-time bank settlement.")
    f2 = doc.add_paragraph()
    add_bullet(f2, "Automated SMS/WhatsApp Alerts: ", "Sending rent due reminders and complaint resolution notifications directly to student and parent phones.")
    f3 = doc.add_paragraph()
    add_bullet(f3, "Biometric & RFID Attendance Integration: ", "Connecting physical turnstile scanners with digital daily in/out logs.")

    # Save document
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")

if __name__ == "__main__":
    create_document("c:\\Users\\bharg\\Desktop\\Hostel\\LuxeHostel_Project_Documentation.docx")
